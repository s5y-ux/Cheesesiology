import docx
import pandas as pd
import random

dates = str(input("Date >: "))
name = str(input("Document Name >: "))

workouts = {"Cardiovascular Endurance": ["Swimming", "Jumping Rope", "Tread Mill", "Battle Ropes"],
"Muscular Strength": ["Squat", "Millitary Press", "Bench Press", "Flys", "One Arm Farmers Walks", "Farmers Walks", "Burpees",
"Lat Pulldown", "Pull Ups", "Rows", "Calf Raises"], "Injury Prevention": ["Cords", "Streching", "Dynamic Streching"]}

keys = list(workouts.keys())
Exercise = list()
Exercise_Component = list()
Date = [dates for count in range(6)]
Duration = [30 for count in range(6)]

global_cock = len(workouts["Muscular Strength"])-1

for count in range(6):
    if (count == 0):
        Exercise_Component.append(keys[0])
        Exercise.append(workouts["Cardiovascular Endurance"][random.randint(0,2)])
    elif(count > 0 and count < 5):
        Exercise_Component.append(keys[1])
        excersize = random.randint(0,global_cock)
        Exercise.append(workouts["Muscular Strength"][excersize])
        workouts["Muscular Strength"].pop(excersize)
        global_cock -= 1
    else:
        Exercise_Component.append(keys[2])
        Exercise.append(workouts["Injury Prevention"][random.randint(0,2)])

si_units = {
    "Date": Date, "Exercise": Exercise, "Exercise Component": Exercise_Component, "Duration": Duration
}
df = pd.DataFrame(si_units)

doc = docx.Document()

t = doc.add_table(rows=(df.shape[0] + 1), cols=df.shape[1])

for j in range(df.shape[1]):
    t.cell(0, j).text = df.columns[j]

for i in range(df.shape[0]):
    for j in range(df.shape[1]):
        cell = df.iat[i, j]
        t.cell(i + 1, j).text = str(cell)

doc.save('{0}.docx'.format(name))
